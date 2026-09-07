# -*- coding: utf-8 -*-
"""检查转换生成的 docx：打开校验 + 打印结构汇总。

用法:
    python check_output.py output\\xxx.docx      # 单个文件
    python check_output.py output                # 整个文件夹
"""
import os
import re
import sys

from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph


def _fix_stdout():
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except Exception:
        pass


def _body_paragraphs(doc):
    """所有正文段落（含目录 sdt、表格内的），按文档顺序。"""
    return [Paragraph(el, doc) for el in doc.element.body.iter(qn('w:p'))]


def _count_drawings(doc):
    return sum(1 for el in doc.element.body.iter() if el.tag == qn('w:drawing'))


def check(path):
    print('=' * 70)
    print('文件:', path)
    try:
        d = Document(path)
    except Exception as e:
        print('  [x] 无法打开（可能结构损坏）:', e)
        return
    print('  [ok] 打开正常')

    print(f'  -- 分节: {len(d.sections)} 个')
    for i, s in enumerate(d.sections):
        print(f'      [{i}] start={s.start_type}  页脚独立={"是" if not s.footer.is_linked_to_previous else "继承前节"}')

    if d.tables:
        print('  -- 封面(第一个表格)有文本的单元格:')
        shown = 0
        for r in d.tables[0].rows:
            t = r.cells[0].text.strip().replace('\n', ' / ')
            if t:
                print(f'      {t[:60]}')
                shown += 1
            if shown >= 8:
                break

    heads = [(p.style.name or '', p.text.strip())
             for p in _body_paragraphs(d)
             if (p.style.name or '').lower().startswith('heading')]
    print(f'  -- 标题 {len(heads)} 个:')
    for st, t in heads:
        print(f'      [{st}] {t}')

    tocs = [p.text for p in _body_paragraphs(d) if (p.style.name or '').lower().startswith('toc')]
    print(f'  -- 目录条目 {len(tocs)} 个:')
    for t in tocs:
        print(f'      {t}')

    print(f'  -- 表格 {len(d.tables)} 个:')
    for i, t in enumerate(d.tables):
        print(f'      [{i}] {len(t.rows)}行 x {len(t.columns)}列  首格={t.rows[0].cells[0].text.strip()[:20]!r}')

    print(f'  -- 图片(drawing) {_count_drawings(d)} 处')

    full = '\n'.join(p.text for p in _body_paragraphs(d))
    leftover = re.findall(r'[X×Ｘx]{1,3}分?系统|GJQ-[A-Z]*XXX', full)
    if leftover:
        print('  [!] 疑似未替换的占位符:', sorted(set(leftover)))
    else:
        print('  [ok] 未发现未替换的占位符')


def main(argv):
    _fix_stdout()
    if not argv:
        print(__doc__)
        return 1
    targets = []
    for a in argv:
        if os.path.isdir(a):
            targets += [os.path.join(a, n) for n in sorted(os.listdir(a))
                        if n.lower().endswith('.docx') and not n.startswith('~$')]
        else:
            targets.append(a)
    if not targets:
        print('未找到 .docx 文件')
        return 1
    for t in targets:
        check(t)
    return 0


if __name__ == '__main__':
    sys.exit(main(sys.argv[1:]))
