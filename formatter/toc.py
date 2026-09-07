# -*- coding: utf-8 -*-
"""Rebuild the table of contents as static entries with placeholder page numbers."""
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

from .util import find_style, add_right_tab


def _has_instr_toc(p_el):
    return any('TOC' in (t.text or '') for t in p_el.iter(qn('w:instrText')))


def _is_field_end_only(p_el):
    """A paragraph holding only the TOC field's 'end' (no nested 'begin')."""
    types = [fc.get(qn('w:fldCharType')) for fc in p_el.iter(qn('w:fldChar'))]
    return 'end' in types and 'begin' not in types


def rebuild_toc(doc, entries, cfg):
    """Replace the template's TOC field (and its cached result) with static entries.

    `entries` is a list of (level, number, title).
    """
    paras = [Paragraph(el, doc) for el in doc.element.body.iter(qn('w:p'))]

    begin = end = None
    for i, p in enumerate(paras):
        if begin is None and _has_instr_toc(p._p):
            begin = i
        elif begin is not None and end is None and _is_field_end_only(p._p):
            end = i
            break
    if begin is None:
        return False
    if end is None:
        # fallback: remove every trailing toc-styled paragraph after begin
        for i in range(begin + 1, len(paras)):
            st = (paras[i].style.name or '').lower()
            if st.startswith('toc'):
                end = i
            else:
                break
        if end is None:
            end = begin

    sec = doc.sections[-1]
    usable = int(sec.page_width - sec.left_margin - sec.right_margin)

    anchor = paras[begin]._p
    for level, number, title in entries:
        if level not in cfg.toc_levels:
            continue
        new_p = doc.add_paragraph()
        style = find_style(doc, f'toc {level}')
        if style is not None:
            new_p.style = style
        label = f'{number} {title}' if number else title
        new_p.add_run(label)
        if cfg.toc_tab_leader:
            add_right_tab(new_p, usable, 'dot')
        new_p.add_run('\t' + cfg.toc_placeholder)
        anchor.addprevious(new_p._p)

    for p in paras[begin:end + 1]:
        parent = p._p.getparent()
        if parent is not None:
            parent.remove(p._p)
    return True
