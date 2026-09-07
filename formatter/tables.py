# -*- coding: utf-8 -*-
"""Table copying: preserve grid, borders, merges, widths; rebuild cell content."""
from copy import deepcopy

from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

from .content import copy_runs


def _gridspan(tc):
    tcPr = tc.find(qn('w:tcPr'))
    if tcPr is None:
        return 1
    gs = tcPr.find(qn('w:gridSpan'))
    if gs is None:
        return 1
    try:
        return max(1, int(gs.get(qn('w:val'))))
    except (TypeError, ValueError):
        return 1


def _copy_tblPr(src_el, new_el):
    src_pr = src_el.find(qn('w:tblPr'))
    if src_pr is None:
        return
    new_pr = new_el.find(qn('w:tblPr'))
    copied = deepcopy(src_pr)
    if new_pr is not None:
        new_el.replace(new_pr, copied)
    else:
        new_el.insert(0, copied)


def _copy_tblGrid(src_el, new_el):
    src_g = src_el.find(qn('w:tblGrid'))
    if src_g is None:
        return
    new_g = new_el.find(qn('w:tblGrid'))
    copied = deepcopy(src_g)
    if new_g is not None:
        new_el.replace(new_g, copied)
    else:
        pr = new_el.find(qn('w:tblPr'))
        if pr is not None:
            pr.addnext(copied)
        else:
            new_el.insert(0, copied)


def _copy_cell(src_tc, dst_cell, doc, ctx, cfg):
    dst_tc = dst_cell._tc
    # preserve cell properties (borders, shading, width, merges, vAlign)
    src_tcPr = src_tc.find(qn('w:tcPr'))
    if src_tcPr is not None:
        dst_tcPr = dst_tc.find(qn('w:tcPr'))
        copied = deepcopy(src_tcPr)
        if dst_tcPr is not None:
            dst_tc.replace(dst_tcPr, copied)
        else:
            dst_tc.insert(0, copied)

    # rebuild paragraph content
    for p in list(dst_tc.findall(qn('w:p'))):
        dst_tc.remove(p)
    for src_p_el in src_tc.findall(qn('w:p')):
        src_p = Paragraph(src_p_el, ctx.source_doc)
        dst_p = dst_cell.add_paragraph()
        copy_runs(src_p, dst_p, ctx.source_doc, cfg, ctx.usable_width_emu)
    if not dst_cell.paragraphs:
        dst_cell.add_paragraph()


def copy_table(src_tbl, doc, ctx, cfg):
    """Copy a source table into `doc`, keeping structure and formatting."""
    src_el = src_tbl._tbl
    tblGrid = src_el.find(qn('w:tblGrid'))
    n_cols = len(tblGrid.findall(qn('w:gridCol'))) if tblGrid is not None else max(1, len(src_tbl.columns))
    n_rows = len(src_el.findall(qn('w:tr')))
    n_cols = max(1, n_cols)
    n_rows = max(1, n_rows)

    new_tbl = doc.add_table(rows=n_rows, cols=n_cols)
    if cfg.table_style:
        from .util import find_style
        st = find_style(doc, cfg.table_style)
        if st is not None:
            try:
                new_tbl.style = st
            except Exception:
                pass

    new_el = new_tbl._tbl
    _copy_tblPr(src_el, new_el)
    _copy_tblGrid(src_el, new_el)

    src_trs = src_el.findall(qn('w:tr'))
    for ri, src_tr in enumerate(src_trs):
        dst_tr = new_el.findall(qn('w:tr'))[ri]
        # row height
        src_trPr = src_tr.find(qn('w:trPr'))
        if src_trPr is not None:
            dst_trPr = dst_tr.find(qn('w:trPr'))
            copied = deepcopy(src_trPr)
            if dst_trPr is not None:
                dst_tr.replace(dst_trPr, copied)
            else:
                dst_tr.insert(0, copied)
        # cells, aligned by gridSpan
        dst_i = 0
        for src_tc in src_tr.findall(qn('w:tc')):
            if dst_i >= n_cols:
                break
            dst_cell = new_tbl.cell(ri, dst_i)
            _copy_cell(src_tc, dst_cell, doc, ctx, cfg)
            dst_i += _gridspan(src_tc)

    ctx.stats.tables += 1
    return new_tbl
