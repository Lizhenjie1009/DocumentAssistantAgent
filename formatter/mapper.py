# -*- coding: utf-8 -*-
"""Core conversion: template as base, map source content onto it."""
import os
import re
from dataclasses import dataclass, field

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from .content import copy_runs
from .tables import copy_table
from .toc import rebuild_toc
from .util import detect_heading_level, strip_number, find_style, has_numbering, is_section_break


@dataclass
class Stats:
    headings: int = 0
    paragraphs: int = 0
    tables: int = 0
    images: int = 0


@dataclass
class Ctx:
    source_doc: object
    usable_width_emu: int
    counters: list = field(default_factory=lambda: [0, 0, 0, 0, 0])
    toc_entries: list = field(default_factory=list)
    stats: Stats = field(default_factory=Stats)
    warnings: list = field(default_factory=list)
    seen_heading: bool = False

    def bump(self, level):
        i = level - 1
        if i < len(self.counters):
            self.counters[i] += 1
            for j in range(i + 1, len(self.counters)):
                self.counters[j] = 0
        return '.'.join(str(c) for c in self.counters[:level] if c > 0)


# --------------------------------------------------------------------------- #
# helpers
# --------------------------------------------------------------------------- #
def _last_is_sectpr(body_el):
    children = list(body_el)
    return bool(children) and children[-1].tag == qn('w:sectPr')


def _source_body_start(source):
    """Index (in body children) where the source's real body begins.

    For style-less sources the front matter (title / cover / hand-typed TOC) is
    detected by taking the first heading that is the LAST occurrence of its
    (level, text) — every TOC entry is duplicated by the real body heading.
    """
    paras = []
    for i, child in enumerate(source.element.body.iterchildren()):
        if child.tag != qn('w:p'):
            continue
        p = Paragraph(child, source)
        lvl = detect_heading_level(p)
        if lvl is not None:
            paras.append((i, lvl, strip_number((p.text or '').strip())))
    last = {}
    for i, l, t in paras:
        last[(l, t)] = i
    for i, l, t in paras:
        if last[(l, t)] == i:
            return i
    return 0


def _find_body_start(doc, cfg):
    body = doc.element.body
    for i, child in enumerate(body.iterchildren()):
        if child.tag != qn('w:p'):
            continue
        p = Paragraph(child, doc)
        name = (p.style.name or '') if p.style is not None else ''
        if name == cfg.body_start_style:
            return i
    return None


def _body_style(src_p, doc):
    st = (src_p.style.name or '') if src_p.style is not None else ''
    low = st.lower()
    if '图题' in st:
        return find_style(doc, '图题') or find_style(doc, 'Caption')
    if 'caption' in low or '题注' in st:
        return find_style(doc, 'Caption')
    if 'list' in low or '列项' in st or has_numbering(src_p):
        return find_style(doc, '一级列项')
    return find_style(doc, '文中正文', '正文', 'Normal')


def _map_paragraph(src_p, doc, ctx, cfg):
    level = detect_heading_level(src_p)
    if level is None:
        st = (src_p.style.name or '') if src_p.style is not None else ''
        # skip leading document-title paragraphs (already handled via title fill)
        if not ctx.seen_heading and st in ('Title', '封面名称', '标题'):
            return None

    new_p = doc.add_paragraph()

    if level is not None:
        level = min(level, 5)
        ctx.seen_heading = True
        style = find_style(doc, f'Heading {level}')
        if style is not None:
            new_p.style = style
        text = strip_number(src_p.text) if cfg.numbering == 'auto' else (src_p.text or '').strip()
        new_p.add_run(text)
        number = ctx.bump(level)
        ctx.toc_entries.append((level, number, text))
        ctx.stats.headings += 1
    else:
        style = _body_style(src_p, doc)
        if style is not None:
            new_p.style = style
        has_img = copy_runs(src_p, new_p, ctx.source_doc, cfg, ctx.usable_width_emu)
        if has_img:
            ctx.stats.images += 1
        if has_img and cfg.image_center and not (src_p.text or '').strip():
            new_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        ctx.stats.paragraphs += 1
    return new_p


def _extract_title(source, source_path, cfg):
    if cfg.title_source == 'filename':
        return os.path.splitext(os.path.basename(source_path))[0]
    # 'auto': Title style first, then the first non-empty title-like paragraph,
    # then the filename stem.
    for p in source.paragraphs:
        name = (p.style.name or '') if p.style is not None else ''
        if name == 'Title' or name == '封面名称' or '标题' in name:
            txt = (p.text or '').strip()
            if txt:
                return txt
    for p in source.paragraphs:
        txt = (p.text or '').strip()
        if not txt:
            continue
        if any(k in txt for k in ('项目名称', '系统名称', '编制单位', '编制日期', '文档编号', '密级')):
            continue
        if len(txt) <= 50:
            return txt
    return os.path.splitext(os.path.basename(source_path))[0]


def _iter_front_paragraphs(doc, limit):
    """Yield Paragraph objects (body + table cells) before index `limit`."""
    for child in list(doc.element.body.iterchildren())[:limit]:
        if child.tag == qn('w:p'):
            yield Paragraph(child, doc)
        elif child.tag == qn('w:tbl'):
            for tr in child.findall(qn('w:tr')):
                for tc in tr.findall(qn('w:tc')):
                    for p_el in tc.findall(qn('w:p')):
                        yield Paragraph(p_el, doc)


def _replace_in_paragraph(p, pat_full, pat_bare, title):
    """Replace placeholder matches that may span multiple runs.

    If the placeholder sits at the start of a short title-like paragraph, the
    whole paragraph text is replaced (avoids duplicated doctype leftovers);
    otherwise only the matched span is replaced (e.g. "项目名称：XX系统").
    """
    runs = p.runs
    if not runs:
        return 0
    full = ''.join(r.text for r in runs)
    bare = False
    m = pat_full.search(full)
    if m is None:
        m = pat_bare.match(full)
        bare = m is not None
    if m is None:
        return 0
    stripped = full.strip()
    if (bare or m.start() == 0) and 0 < len(stripped) <= 60:
        runs[0].text = title
        for r in runs[1:]:
            r.text = ''
        return 1
    s, e = m.span()
    new_full = full[:s] + title + full[e:]
    runs[0].text = new_full
    for r in runs[1:]:
        r.text = ''
    return 1


def _fill_title(doc, source, source_path, cfg, ctx):
    if not cfg.title_source or not cfg.title_placeholder_regex:
        return
    title = _extract_title(source, source_path, cfg)
    if not title:
        return
    pat_full = re.compile(cfg.title_placeholder_regex)
    pat_bare = re.compile(r'^[X×Ｘx]{1,3}(?=\s|$)')
    body_start = _find_body_start(doc, cfg)
    limit = body_start if body_start is not None else len(list(doc.element.body))
    count = 0
    for p in _iter_front_paragraphs(doc, limit):
        count += _replace_in_paragraph(p, pat_full, pat_bare, title)
    if count:
        ctx.warnings.append(f'已把 {count} 处标题占位符替换为“{title}”')


# --------------------------------------------------------------------------- #
# main entry
# --------------------------------------------------------------------------- #
def convert(source_path, tpl_cfg, output_path):
    source = Document(source_path)
    template = Document(tpl_cfg.file)
    cfg = tpl_cfg

    sec = template.sections[-1]
    usable = int(sec.page_width - sec.left_margin - sec.right_margin)
    ctx = Ctx(source_doc=source, usable_width_emu=usable)

    body_start = _find_body_start(template, cfg)
    if body_start is None:
        ctx.warnings.append('模板中未找到正文起点样式 %r，将把源内容追加到模板末尾' % cfg.body_start_style)
        body_start = len(list(template.element.body))

    body_el = template.element.body
    children = list(body_el)
    end = len(children) - 1 if _last_is_sectpr(body_el) else len(children)
    for child in children[body_start:end]:
        body_el.remove(child)

    # map source body (paragraphs and tables; skip front matter, section breaks & sectPr)
    src_start = _source_body_start(source)
    for idx, child in enumerate(source.element.body.iterchildren()):
        if idx < src_start:
            continue
        tag = child.tag
        if tag == qn('w:p'):
            p = Paragraph(child, source)
            if is_section_break(p):
                continue
            _map_paragraph(p, template, ctx, cfg)
        elif tag == qn('w:tbl'):
            t = Table(child, source)
            copy_table(t, template, ctx, cfg)

    rebuild_toc(template, ctx.toc_entries, cfg)
    _fill_title(template, source, source_path, cfg, ctx)

    out_dir = os.path.dirname(output_path)
    if out_dir:
        os.makedirs(out_dir, exist_ok=True)
    template.save(output_path)
    return ctx
